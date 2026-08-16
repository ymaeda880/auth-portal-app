# -*- coding: utf-8 -*-
# auth_portal_app/lib/sks_library/report_builders.py
# ============================================================
# SKSライブラリー 検索結果出力
#
# 機能：
# - 検索結果をText形式で生成する
# - 検索結果をWord形式で生成する
# - 検索結果をPDF形式で生成する
# - 検索結果をExcel形式で生成する
#
# 方針：
# - 検索結果DataFrameを共通入力として使用する
# - キーワード検索とAI検索の両方に対応する
# - AI検索の場合はAI関連度・AI選定理由も出力する
# - 元のSKS_booklist.csvは変更しない
# ============================================================

from __future__ import annotations

# ============================================================
# imports
# ============================================================
import datetime as _dt
import html
import io
from typing import Any

import pandas as pd


# ============================================================
# 定数
# ============================================================

# ------------------------------------------------------------
# 内部ID
# ------------------------------------------------------------
INTERNAL_ID_COLUMN = "_sks_id"

# ------------------------------------------------------------
# 通常の書誌情報
# ------------------------------------------------------------
BOOK_COLUMNS = [
    "分類: NDC",
    "棚番号-段",
    "タイトル",
    "編著者",
    "発行者",
    "発行年",
    "言語",
    "注記：原著名／仮訳、ISBN等",
    "貴重書",
    "タグ",
]

# ------------------------------------------------------------
# AI検索時の追加情報
# ------------------------------------------------------------
AI_COLUMNS = [
    "AI関連度",
    "AI選定理由",
]


# ============================================================
# helper：値を安全な文字列へ
# ============================================================
def _as_text(
    value: Any,
) -> str:
    if value is None:
        return ""

    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass

    return str(value).strip()


# ============================================================
# helper：AI検索結果か
# ============================================================
def _is_ai_result(
    df: pd.DataFrame,
) -> bool:
    return (
        "AI関連度" in df.columns
        or "AI選定理由" in df.columns
    )


# ============================================================
# helper：出力列
# ============================================================
def _resolve_output_columns(
    df: pd.DataFrame,
) -> list[str]:

    columns = [
        col
        for col in BOOK_COLUMNS
        if col in df.columns
    ]

    if _is_ai_result(df):
        for col in AI_COLUMNS:
            if col in df.columns:
                columns.append(col)

    return columns


# ============================================================
# helper：検索種別ラベル
# ============================================================
def _search_mode_label(
    search_mode: str,
) -> str:

    raw = str(
        search_mode or ""
    ).strip()

    if raw.lower() in {
        "ai",
        "ai検索",
    }:
        return "AI検索"

    return "キーワード検索"


# ============================================================
# helper：検索条件テキスト
# ============================================================
def _build_search_condition_lines(
    *,
    search_mode: str,
    query: str,
    keyword_mode: str = "",
    keyword_target: str = "",
    model: str = "",
) -> list[str]:

    lines = [
        f"検索方法: {_search_mode_label(search_mode)}",
        f"検索内容: {str(query or '').strip()}",
    ]

    if _search_mode_label(search_mode) == "キーワード検索":

        if str(keyword_mode or "").strip():
            lines.append(
                f"検索条件: {str(keyword_mode).strip()}"
            )

        if str(keyword_target or "").strip():
            lines.append(
                f"検索対象: {str(keyword_target).strip()}"
            )

    else:

        if str(model or "").strip():
            lines.append(
                f"使用モデル: {str(model).strip()}"
            )

    return lines


# ============================================================
# helper：1冊分のText
# ============================================================
def _build_book_text(
    row: pd.Series,
    *,
    number: int,
    include_ai: bool,
) -> str:

    lines: list[str] = []

    # ------------------------------------------------------------
    # 番号・タイトル
    # ------------------------------------------------------------
    title = _as_text(
        row.get(
            "タイトル",
            "",
        )
    )

    lines.append(
        f"{number}. {title or '(タイトルなし)'}"
    )

    # ------------------------------------------------------------
    # 書誌情報
    # ------------------------------------------------------------
    metadata_columns = [
        ("分類", "分類: NDC"),
        ("棚番号", "棚番号-段"),
        ("編著者", "編著者"),
        ("発行者", "発行者"),
        ("発行年", "発行年"),
        ("言語", "言語"),
        ("注記", "注記：原著名／仮訳、ISBN等"),
        ("貴重書", "貴重書"),
        ("タグ", "タグ"),
    ]

    for label, col in metadata_columns:

        if col not in row.index:
            continue

        value = _as_text(
            row.get(
                col,
                "",
            )
        )

        if not value:
            continue

        lines.append(
            f"{label}: {value}"
        )

    # ------------------------------------------------------------
    # AI情報
    # ------------------------------------------------------------
    if include_ai:

        score = _as_text(
            row.get(
                "AI関連度",
                "",
            )
        )

        reason = _as_text(
            row.get(
                "AI選定理由",
                "",
            )
        )

        if score:
            lines.append("")
            lines.append(
                f"AI関連度: {score}"
            )

        if reason:
            lines.append("")
            lines.append(
                "AIが選択した理由:"
            )
            lines.append(
                reason
            )

    return "\n".join(
        lines
    )


# ============================================================
# Text 出力
# ============================================================
def build_sks_search_txt_bytes(
    *,
    result_df: pd.DataFrame,
    search_mode: str,
    query: str,
    keyword_mode: str = "",
    keyword_target: str = "",
    model: str = "",
) -> bytes:

    # ------------------------------------------------------------
    # ヘッダー
    # ------------------------------------------------------------
    sections: list[str] = [
        "=== SKSライブラリー検索結果 ===",
        f"生成日時: {_dt.datetime.now():%Y-%m-%d %H:%M:%S}",
    ]

    sections.extend(
        _build_search_condition_lines(
            search_mode=search_mode,
            query=query,
            keyword_mode=keyword_mode,
            keyword_target=keyword_target,
            model=model,
        )
    )

    sections.extend(
        [
            f"検索結果: {len(result_df)}冊",
            "",
        ]
    )

    # ------------------------------------------------------------
    # 検索結果
    # ------------------------------------------------------------
    if result_df.empty:

        sections.append(
            "該当する蔵書はありませんでした．"
        )

    else:

        include_ai = _is_ai_result(
            result_df
        )

        for i, (_, row) in enumerate(
            result_df.iterrows(),
            start=1,
        ):

            if i > 1:
                sections.extend(
                    [
                        "",
                        "----------------------------------------",
                        "",
                    ]
                )

            sections.append(
                _build_book_text(
                    row,
                    number=i,
                    include_ai=include_ai,
                )
            )

    data = "\n".join(
        sections
    )

    return data.encode(
        "utf-8"
    )


# ============================================================
# Word 出力
# ============================================================
def build_sks_search_docx_bytes(
    *,
    result_df: pd.DataFrame,
    search_mode: str,
    query: str,
    keyword_mode: str = "",
    keyword_target: str = "",
    model: str = "",
) -> tuple[bytes, str]:

    # ------------------------------------------------------------
    # import
    # ------------------------------------------------------------
    try:
        from docx import Document
        from docx.shared import Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH

    except Exception:

        return (
            build_sks_search_txt_bytes(
                result_df=result_df,
                search_mode=search_mode,
                query=query,
                keyword_mode=keyword_mode,
                keyword_target=keyword_target,
                model=model,
            ),
            ".txt",
        )

    # ------------------------------------------------------------
    # document
    # ------------------------------------------------------------
    doc = Document()

    # ------------------------------------------------------------
    # title
    # ------------------------------------------------------------
    title = doc.add_heading(
        "SKSライブラリー検索結果",
        0,
    )

    try:
        title.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )
    except Exception:
        pass

    # ------------------------------------------------------------
    # metadata
    # ------------------------------------------------------------
    p = doc.add_paragraph()

    p.add_run(
        "生成日時: "
    ).bold = True

    p.add_run(
        _dt.datetime.now().strftime(
            "%Y-%m-%d %H:%M:%S"
        )
    )

    condition_lines = (
        _build_search_condition_lines(
            search_mode=search_mode,
            query=query,
            keyword_mode=keyword_mode,
            keyword_target=keyword_target,
            model=model,
        )
    )

    for line in condition_lines:

        if ": " in line:
            label, value = line.split(
                ": ",
                1,
            )

            p.add_run(
                f"\n{label}: "
            ).bold = True

            p.add_run(
                value
            )

        else:
            p.add_run(
                f"\n{line}"
            )

    p.add_run(
        "\n検索結果: "
    ).bold = True

    p.add_run(
        f"{len(result_df)}冊"
    )

    # ------------------------------------------------------------
    # 結果なし
    # ------------------------------------------------------------
    if result_df.empty:

        doc.add_paragraph(
            "該当する蔵書はありませんでした．"
        )

    # ------------------------------------------------------------
    # 結果あり
    # ------------------------------------------------------------
    else:

        include_ai = _is_ai_result(
            result_df
        )

        for i, (_, row) in enumerate(
            result_df.iterrows(),
            start=1,
        ):

            book_title = _as_text(
                row.get(
                    "タイトル",
                    "",
                )
            )

            doc.add_heading(
                f"{i}. {book_title or '(タイトルなし)'}",
                level=1,
            )

            # --------------------------------------------------------
            # 書誌情報
            # --------------------------------------------------------
            metadata_columns = [
                ("分類", "分類: NDC"),
                ("棚番号", "棚番号-段"),
                ("編著者", "編著者"),
                ("発行者", "発行者"),
                ("発行年", "発行年"),
                ("言語", "言語"),
                (
                    "注記",
                    "注記：原著名／仮訳、ISBN等",
                ),
                ("貴重書", "貴重書"),
                ("タグ", "タグ"),
            ]

            p_book = doc.add_paragraph()

            first_line = True

            for label, col in metadata_columns:

                if col not in result_df.columns:
                    continue

                value = _as_text(
                    row.get(
                        col,
                        "",
                    )
                )

                if not value:
                    continue

                if not first_line:
                    p_book.add_run(
                        "\n"
                    )

                run_label = p_book.add_run(
                    f"{label}: "
                )

                run_label.bold = True

                p_book.add_run(
                    value
                )

                first_line = False

            # --------------------------------------------------------
            # AI情報
            # --------------------------------------------------------
            if include_ai:

                score = _as_text(
                    row.get(
                        "AI関連度",
                        "",
                    )
                )

                reason = _as_text(
                    row.get(
                        "AI選定理由",
                        "",
                    )
                )

                if score:

                    p_score = doc.add_paragraph()

                    p_score.add_run(
                        "AI関連度: "
                    ).bold = True

                    p_score.add_run(
                        score
                    )

                if reason:

                    p_reason = doc.add_paragraph()

                    p_reason.add_run(
                        "AIが選択した理由:"
                    ).bold = True

                    p_reason.add_run(
                        "\n"
                    )

                    r = p_reason.add_run(
                        reason
                    )

                    try:
                        r.font.size = Pt(10)
                    except Exception:
                        pass

    # ------------------------------------------------------------
    # bytes
    # ------------------------------------------------------------
    bio = io.BytesIO()

    doc.save(
        bio
    )

    return (
        bio.getvalue(),
        ".docx",
    )


# ============================================================
# PDF 出力
# ============================================================
def build_sks_search_pdf_bytes(
    *,
    result_df: pd.DataFrame,
    search_mode: str,
    query: str,
    keyword_mode: str = "",
    keyword_target: str = "",
    model: str = "",
) -> bytes:

    # ------------------------------------------------------------
    # import
    # ------------------------------------------------------------
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import (
            TA_CENTER,
            TA_LEFT,
        )
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import (
            ParagraphStyle,
            getSampleStyleSheet,
        )
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.platypus import (
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )

    except Exception:
        return b""

    # ------------------------------------------------------------
    # Japanese font
    # ------------------------------------------------------------
    font_name = None

    for fname in (
        "HeiseiMin-W3",
        "HeiseiKakuGo-W5",
    ):

        try:
            pdfmetrics.registerFont(
                UnicodeCIDFont(
                    fname
                )
            )

            font_name = fname
            break

        except Exception:
            continue

    if not font_name:
        return b""

    # ------------------------------------------------------------
    # document
    # ------------------------------------------------------------
    buf = io.BytesIO()

    pagesize = A4

    doc = SimpleDocTemplate(
        buf,
        pagesize=pagesize,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title="SKSライブラリー検索結果",
    )

    # ------------------------------------------------------------
    # styles
    # ------------------------------------------------------------
    styles = getSampleStyleSheet()

    styles.add(
        ParagraphStyle(
            name="SKSTitle",
            fontName=font_name,
            fontSize=18,
            leading=22,
            alignment=TA_CENTER,
            spaceAfter=8,
        )
    )

    styles.add(
        ParagraphStyle(
            name="SKSMeta",
            fontName=font_name,
            fontSize=9.5,
            leading=13,
            alignment=TA_LEFT,
            spaceAfter=3,
            wordWrap="CJK",
        )
    )

    styles.add(
        ParagraphStyle(
            name="SKSBookTitle",
            fontName=font_name,
            fontSize=12,
            leading=16,
            spaceBefore=8,
            spaceAfter=5,
            wordWrap="CJK",
        )
    )

    styles.add(
        ParagraphStyle(
            name="SKSBody",
            fontName=font_name,
            fontSize=9.5,
            leading=13,
            alignment=TA_LEFT,
            wordWrap="CJK",
        )
    )

    styles.add(
        ParagraphStyle(
            name="SKSAIReason",
            fontName=font_name,
            fontSize=9.5,
            leading=14,
            alignment=TA_LEFT,
            leftIndent=4 * mm,
            rightIndent=4 * mm,
            wordWrap="CJK",
        )
    )

    # ------------------------------------------------------------
    # page number
    # ------------------------------------------------------------
    def _page_number(
        canvas,
        doc_,
    ) -> None:

        canvas.setFont(
            font_name,
            9,
        )

        canvas.drawRightString(
            pagesize[0] - doc_.rightMargin,
            10 * mm,
            f"{doc_.page}",
        )

    # ------------------------------------------------------------
    # story
    # ------------------------------------------------------------
    story: list[Any] = []

    story.append(
        Paragraph(
            "SKSライブラリー検索結果",
            styles["SKSTitle"],
        )
    )

    # ------------------------------------------------------------
    # metadata
    # ------------------------------------------------------------
    meta_lines = [
        f"生成日時：{_dt.datetime.now():%Y-%m-%d %H:%M:%S}",
    ]

    meta_lines.extend(
        [
            line.replace(
                ": ",
                "：",
                1,
            )
            for line in _build_search_condition_lines(
                search_mode=search_mode,
                query=query,
                keyword_mode=keyword_mode,
                keyword_target=keyword_target,
                model=model,
            )
        ]
    )

    meta_lines.append(
        f"検索結果：{len(result_df)}冊"
    )

    for line in meta_lines:

        story.append(
            Paragraph(
                html.escape(
                    line
                ),
                styles["SKSMeta"],
            )
        )

    story.append(
        Spacer(
            1,
            5 * mm,
        )
    )

    # ------------------------------------------------------------
    # 結果なし
    # ------------------------------------------------------------
    if result_df.empty:

        story.append(
            Paragraph(
                "該当する蔵書はありませんでした．",
                styles["SKSBody"],
            )
        )

    # ------------------------------------------------------------
    # 結果あり
    # ------------------------------------------------------------
    else:

        include_ai = _is_ai_result(
            result_df
        )

        for i, (_, row) in enumerate(
            result_df.iterrows(),
            start=1,
        ):

            title = _as_text(
                row.get(
                    "タイトル",
                    "",
                )
            )

            story.append(
                Paragraph(
                    html.escape(
                        f"{i}. {title or '(タイトルなし)'}"
                    ),
                    styles["SKSBookTitle"],
                )
            )

            # --------------------------------------------------------
            # 書誌情報テーブル
            # --------------------------------------------------------
            metadata_columns = [
                ("分類", "分類: NDC"),
                ("棚番号", "棚番号-段"),
                ("編著者", "編著者"),
                ("発行者", "発行者"),
                ("発行年", "発行年"),
                ("言語", "言語"),
                (
                    "注記",
                    "注記：原著名／仮訳、ISBN等",
                ),
                ("貴重書", "貴重書"),
                ("タグ", "タグ"),
            ]

            rows = []

            for label, col in metadata_columns:

                if col not in result_df.columns:
                    continue

                value = _as_text(
                    row.get(
                        col,
                        "",
                    )
                )

                if not value:
                    continue

                rows.append(
                    [
                        Paragraph(
                            html.escape(
                                label
                            ),
                            styles["SKSBody"],
                        ),
                        Paragraph(
                            html.escape(
                                value
                            ).replace(
                                "\n",
                                "<br/>",
                            ),
                            styles["SKSBody"],
                        ),
                    ]
                )

            if rows:

                table = Table(
                    rows,
                    colWidths=[
                        25 * mm,
                        130 * mm,
                    ],
                    splitByRow=1,
                )

                table.setStyle(
                    TableStyle(
                        [
                            (
                                "VALIGN",
                                (0, 0),
                                (-1, -1),
                                "TOP",
                            ),
                            (
                                "LEFTPADDING",
                                (0, 0),
                                (-1, -1),
                                4,
                            ),
                            (
                                "RIGHTPADDING",
                                (0, 0),
                                (-1, -1),
                                4,
                            ),
                            (
                                "TOPPADDING",
                                (0, 0),
                                (-1, -1),
                                3,
                            ),
                            (
                                "BOTTOMPADDING",
                                (0, 0),
                                (-1, -1),
                                3,
                            ),
                            (
                                "LINEBELOW",
                                (0, 0),
                                (-1, -1),
                                0.2,
                                colors.lightgrey,
                            ),
                        ]
                    )
                )

                story.append(
                    table
                )

            # --------------------------------------------------------
            # AI情報
            # --------------------------------------------------------
            if include_ai:

                score = _as_text(
                    row.get(
                        "AI関連度",
                        "",
                    )
                )

                reason = _as_text(
                    row.get(
                        "AI選定理由",
                        "",
                    )
                )

                if score:

                    story.append(
                        Spacer(
                            1,
                            3 * mm,
                        )
                    )

                    story.append(
                        Paragraph(
                            f"<b>AI関連度：</b>{html.escape(score)}",
                            styles["SKSBody"],
                        )
                    )

                if reason:

                    story.append(
                        Spacer(
                            1,
                            2 * mm,
                        )
                    )

                    story.append(
                        Paragraph(
                            "<b>AIが選択した理由：</b>",
                            styles["SKSBody"],
                        )
                    )

                    story.append(
                        Paragraph(
                            html.escape(
                                reason
                            ).replace(
                                "\n",
                                "<br/>",
                            ),
                            styles["SKSAIReason"],
                        )
                    )

            story.append(
                Spacer(
                    1,
                    7 * mm,
                )
            )

    # ------------------------------------------------------------
    # build
    # ------------------------------------------------------------
    doc.build(
        story,
        onFirstPage=_page_number,
        onLaterPages=_page_number,
    )

    pdf_bytes = buf.getvalue()

    buf.close()

    return pdf_bytes


# ============================================================
# Excel 出力
# ============================================================
def build_sks_search_xlsx_bytes(
    *,
    result_df: pd.DataFrame,
    search_mode: str,
    query: str,
    keyword_mode: str = "",
    keyword_target: str = "",
    model: str = "",
) -> bytes:

    # ------------------------------------------------------------
    # import
    # ------------------------------------------------------------
    try:
        from openpyxl import Workbook
        from openpyxl.styles import (
            Alignment,
            Font,
        )
        from openpyxl.utils import get_column_letter

    except Exception as e:
        raise RuntimeError(
            "Excel出力には openpyxl が必要です．"
        ) from e

    # ------------------------------------------------------------
    # workbook
    # ------------------------------------------------------------
    wb = Workbook()

    ws = wb.active

    ws.title = "検索結果"

    # ------------------------------------------------------------
    # 検索情報
    # ------------------------------------------------------------
    ws["A1"] = "SKSライブラリー検索結果"
    ws["A1"].font = Font(
        bold=True,
        size=14,
    )

    ws["A2"] = "生成日時"
    ws["B2"] = _dt.datetime.now().strftime(
        "%Y-%m-%d %H:%M:%S"
    )

    condition_lines = (
        _build_search_condition_lines(
            search_mode=search_mode,
            query=query,
            keyword_mode=keyword_mode,
            keyword_target=keyword_target,
            model=model,
        )
    )

    row_index = 3

    for line in condition_lines:

        if ": " in line:

            label, value = line.split(
                ": ",
                1,
            )

            ws.cell(
                row=row_index,
                column=1,
                value=label,
            )

            ws.cell(
                row=row_index,
                column=2,
                value=value,
            )

        else:

            ws.cell(
                row=row_index,
                column=1,
                value=line,
            )

        row_index += 1

    ws.cell(
        row=row_index,
        column=1,
        value="検索結果",
    )

    ws.cell(
        row=row_index,
        column=2,
        value=f"{len(result_df)}冊",
    )

    row_index += 2

    # ------------------------------------------------------------
    # 列
    # ------------------------------------------------------------
    output_columns = (
        _resolve_output_columns(
            result_df
        )
    )

    # ------------------------------------------------------------
    # ヘッダー
    # ------------------------------------------------------------
    header_row = row_index

    for col_index, col_name in enumerate(
        output_columns,
        start=1,
    ):

        cell = ws.cell(
            row=header_row,
            column=col_index,
            value=col_name,
        )

        cell.font = Font(
            bold=True,
        )

        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

    # ------------------------------------------------------------
    # データ
    # ------------------------------------------------------------
    for data_row_index, (_, row) in enumerate(
        result_df.iterrows(),
        start=header_row + 1,
    ):

        for col_index, col_name in enumerate(
            output_columns,
            start=1,
        ):

            value = _as_text(
                row.get(
                    col_name,
                    "",
                )
            )

            # --------------------------------------------------------
            # すべて文字列として保存
            #
            # NDC・棚番号・発行年などのExcel自動変換を防ぐ．
            # --------------------------------------------------------
            cell = ws.cell(
                row=data_row_index,
                column=col_index,
                value=value,
            )

            cell.number_format = "@"

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    # ------------------------------------------------------------
    # フリーズ
    # ------------------------------------------------------------
    ws.freeze_panes = (
        f"A{header_row + 1}"
    )

    # ------------------------------------------------------------
    # オートフィルター
    # ------------------------------------------------------------
    if output_columns:

        last_col = get_column_letter(
            len(output_columns)
        )

        last_row = max(
            header_row,
            header_row + len(result_df),
        )

        ws.auto_filter.ref = (
            f"A{header_row}:{last_col}{last_row}"
        )

    # ------------------------------------------------------------
    # 列幅
    # ------------------------------------------------------------
    preferred_widths = {
        "分類: NDC": 14,
        "棚番号-段": 12,
        "タイトル": 45,
        "編著者": 28,
        "発行者": 28,
        "発行年": 18,
        "言語": 12,
        "注記：原著名／仮訳、ISBN等": 45,
        "貴重書": 12,
        "タグ": 25,
        "AI関連度": 12,
        "AI選定理由": 55,
    }

    for index, col_name in enumerate(
        output_columns,
        start=1,
    ):

        width = preferred_widths.get(
            col_name,
            18,
        )

        ws.column_dimensions[
            get_column_letter(
                index
            )
        ].width = width

    # ------------------------------------------------------------
    # 検索条件部の列幅
    # ------------------------------------------------------------
    if ws.column_dimensions["A"].width < 18:
        ws.column_dimensions["A"].width = 18

    # ------------------------------------------------------------
    # bytes
    # ------------------------------------------------------------
    bio = io.BytesIO()

    wb.save(
        bio
    )

    return bio.getvalue()